#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Compose the final MOH-branded guide from content_model.json (the
mechanically-extracted, complete ground truth of source.docx)."""
import json
import re
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, '/home/user/mudiu-website/docwork/build')
from moh_style import (
    GREEN, GREEN_DARK, GREEN_TINT, GREEN_TINT2, GOLD, GOLD_TINT, CREAM, BROWN,
    WHITE, GRAY_TEXT, GRAY_LIGHT, GRAY_MED, RED, RED_TINT, FONT_BODY, FONT_HEAD,
    rtl_paragraph, set_font, add_text, shade_paragraph, border_paragraph,
    paragraph_spacing, paragraph_indent, keep_with_next, keep_lines,
    page_break_before, add_page_break, add_page_number, add_toc,
    set_table_rtl, shade_cell, cell_borders, cell_margins, cell_valign,
    set_col_widths, write_cell, style_table_default, style_scale_table,
    add_picture_centered, gold_dot, hr_rule, spacer,
)
from components import (
    setup_section, setup_styles, build_header, build_footer, callout,
    principle_card, definition_box, example_box, process_flow,
    process_flow_horizontal, concept_cards_grid, classification_cards,
    output_cards, style_role_table, style_log_table, force_update_fields,
    LOGO, BANNER, PAGE_W, PAGE_H, MARGIN,
)

MODEL_PATH = "/home/user/mudiu-website/docwork/build/content_model.json"
OUT_PATH = "/home/user/mudiu-website/docwork/build/final_guide.docx"

DOC_TITLE = "الدليل الإرشادي لتطبيق منهجيتي المؤشرات والمشاريع"
DOC_SUB = "من التصنيف إلى القرار: الأساس المرجعي لترتيب أولوية المؤشرات والمشاريع"

blocks = json.load(open(MODEL_PATH, encoding="utf-8"))
B = {b["idx"]: b for b in blocks}
N = len(blocks)


# ============================================================== helpers ==

def clean(text):
    """Light typographic normalization — never changes wording, only
    spacing/punctuation consistency (per the language-review pass)."""
    t = text.strip()
    t = re.sub(r'^\.\s*', '', t)            # stray leading '.' e.g. '.C1 ...'
    t = re.sub(r'^\s*-\s*', '', t)           # stray leading '-'
    t = re.sub(r'\s+', ' ', t)
    t = re.sub(r'\s*%', '%', t)              # '20 %' / '20  %' -> '20%'
    t = re.sub(r'--\s*', '– ', t)            # '--' -> en dash
    t = re.sub(r'\s*–\s*', ' – ', t)
    t = re.sub(r'(\d)\.(\d)\s+', r'\1.\2 ', t)  # '1.1  ' -> '1.1 '
    t = re.sub(r'([ء-ي])(\d)', r'\1 \2', t)  # 'الاستراتيجية20%' -> '... 20%'
    return t.strip()


def is_flow_text(text):
    return '\n↓\n' in text or bool(re.search(r'\s→\s', text))


def split_flow(text):
    if '\n↓\n' in text:
        return 'v', [s.strip() for s in text.split('\n↓\n')]
    return 'h', [s.strip() for s in re.split(r'\s→\s', text)]


BOLD_PRINCIPLE_RE = re.compile(r'^[^\n]{6,200}[.؟]$')


def is_principle_paragraph(b):
    if b["kind"] != "p" or b["is_list"] or not b["text"].strip():
        return False
    runs = [r for r in b["runs"] if r["t"].strip()]
    if not runs or not all(r["b"] for r in runs):
        return False
    t = b["text"].strip()
    if "\n" in t or len(t) > 220 or len(t) < 8:
        return False
    if t.startswith(("القاعدة", "المرحلة", "الجزء", "قاعدة")):
        return False
    return True


# ===================================================== structural maps ==

HEADING_LEVEL_FIX = {}
for i in (490, 508, 521, 534, 547, 559, 570, 578, 595, 610, 638, 649, 668,
          849, 904, 925, 765, 817):
    HEADING_LEVEL_FIX[i] = 2
for i in (767, 787, 819, 837, 851, 868, 888, 905, 926):
    HEADING_LEVEL_FIX[i] = 3
for i in (768, 770, 772, 778, 788, 790, 792, 804, 809, 820, 823, 832, 838,
          840, 842, 852, 861, 863, 869, 871, 882, 889, 900, 906, 908, 916,
          927, 929, 941, 479, 494, 510, 480, 481, 482, 483, 495, 496, 497,
          498, 511, 512, 513, 514):
    HEADING_LEVEL_FIX[i] = 4

FORCE_HEADING = {213: 3, 218: 3, 261: 3, 269: 3, 277: 3, 1244: 2}
PRINCIPLE_EXCLUDE = {1245}

# idx -> (kicker, chapter-purpose idx to quote a one-line teaser from)
CHAPTER_STARTS = {
    52: ("مقدمة الدليل", 53),
    104: ("الفصل الأول", 106),
    295: ("الفصل الثاني", 296),
    415: ("الفصل الثالث", 417),
    714: ("الفصل الرابع", 716),
    1081: ("الفصل الخامس", 1083),
}
FRONT_EYEBROW_IDX = 39  # 'الإطار العام للمنهجيات' — folded into the first divider

CASE_RANGES = [(219, 222), (223, 228), (229, 234), (235, 238), (240, 244),
               (245, 251), (252, 260)]
CASE_START_IDX = {r[0] for r in CASE_RANGES}
CASE_SKIP = set()
for s, e in CASE_RANGES:
    CASE_SKIP.update(range(s, e + 1))

SIMPLE_EXAMPLE_RANGE = (118, 138)

CLASS_LEGEND_TRIGGER = 263
CLASS_LEGEND_SKIP = {263, 264, 265, 266}

FLOW_TRIGGER = 279
FLOW_SKIP = set(range(279, 290))

GATE_INSERTS = {
    486: ('flow', ["الارتباط الاستراتيجي", "صحة القياس", "قيمة النتيجة للمستفيد",
                   "قيمة المؤشر في القرار"]),
    503: ('grid', ["الأولوية الصحية للنتيجة", "قيمة التغيير المتحقق",
                   "قدرة الوزارة على التأثير", "واقعية ظهور النتيجة زمنيًا"]),
    517: ('grid', ["اكتمال بطاقة المؤشر", "ملاءمته للهدف الاستراتيجي",
                   "توفر بياناته بانتظام", "نضجه للاستخدام الاستراتيجي"]),
}

DECISION_MATRIX_IDX = 1222  # 'ج. مصفوفة تتبع القرار' -> render as flow, not '→' text
ESCALATION_AFTER_IDX = 1176  # after 'الجهة المخولة باتخاذ القرار' role block

TABLE_STYLE_BY_ORDER = {
    0: 'role2col', 1: 'default', 2: 'default', 3: 'default', 4: 'default',
    5: 'weight3', 6: 'cf_scale', 7: 'weight3', 8: 'scale', 9: 'scale',
    10: 'scale', 11: 'scale', 12: 'scale', 13: 'scale', 14: 'scale',
    15: 'scale', 16: 'scale', 17: 'weight2', 18: 'log', 19: 'log',
}

DEFINITION_HEADINGS = {'تعريف الاتساق الاستراتيجي'}


# ============================================================ builders ==

def golden_ribbon(doc, current):
    steps = ["الاتساق الاستراتيجي", "المنهجيات", "التطبيق والحوكمة"]
    tbl = doc.add_table(rows=1, cols=len(steps))
    set_table_rtl(tbl)
    tbl.autofit = False
    set_col_widths(tbl, [15.5 / len(steps)] * len(steps))
    for i, s in enumerate(steps):
        cell = tbl.rows[0].cells[i]
        active = (s == current)
        p = cell.paragraphs[0]
        rtl_paragraph(p, align='center')
        paragraph_spacing(p, before=3, after=3)
        add_text(p, s, name=FONT_HEAD, size=8.7, bold=active,
                 color=(WHITE if active else "d9e4dd"))
        shade_cell(cell, GREEN if active else GREEN_DARK)
        cell_borders(cell, color=WHITE, sz=2)
        cell_valign(cell, 'center')
        cell_margins(cell, top=30, bottom=30, start=40, end=40)
    spacer(doc, 2)


def _blank(doc, h=10):
    p = doc.add_paragraph()
    paragraph_spacing(p, before=0, after=h)
    return p


def cover_page(doc):
    _blank(doc, 16)
    add_picture_centered(doc, LOGO, width_cm=2.9)
    _blank(doc, 22)

    p = doc.add_paragraph()
    rtl_paragraph(p, align='center')
    shade_paragraph(p, GREEN)
    paragraph_spacing(p, before=16, after=4)
    add_text(p, "مكتب إدارة الاستراتيجية", name=FONT_HEAD, size=12, bold=True,
             color=CREAM)

    p2 = doc.add_paragraph()
    rtl_paragraph(p2, align='center')
    shade_paragraph(p2, GREEN)
    paragraph_spacing(p2, before=4, after=4, line=1.25)
    add_text(p2, DOC_TITLE, name=FONT_HEAD, size=25, bold=True, color=WHITE)

    p3 = doc.add_paragraph()
    rtl_paragraph(p3, align='center')
    shade_paragraph(p3, GREEN)
    paragraph_spacing(p3, before=4, after=16, line=1.25)
    add_text(p3, DOC_SUB, name=FONT_BODY, size=11.5, color="eef3ee")

    _blank(doc, 18)

    p4 = doc.add_paragraph()
    rtl_paragraph(p4, align='center')
    border_paragraph(p4, color=GOLD, sz=8, sides=('top', 'bottom'), space=6)
    paragraph_spacing(p4, before=6, after=6)
    add_text(p4, "مسودة للمراجعة والاعتماد", name=FONT_HEAD, size=12, bold=True,
             color=BROWN)

    p5 = doc.add_paragraph()
    rtl_paragraph(p5, align='center')
    paragraph_spacing(p5, before=8, after=2)
    add_text(p5, "لا تمثل هذه النسخة وثيقة معتمدة، وتبقى جميع القواعد الواردة "
                 "فيها خاضعة لإجراءات الاعتماد الرسمية", name=FONT_BODY, size=9.5,
             italic=True, color=GRAY_TEXT)
    p6 = doc.add_paragraph()
    rtl_paragraph(p6, align='center')
    paragraph_spacing(p6, before=0, after=4)
    add_text(p6, "سري – للاستخدام الداخلي", name=FONT_BODY, size=9.5, italic=True,
             color=GRAY_TEXT)

    _blank(doc, 30)
    gold_dot(doc)


def toc_page(doc):
    h = doc.add_paragraph()
    rtl_paragraph(h)
    add_text(h, "جدول المحتويات", name=FONT_HEAD, size=20, bold=True, color=GREEN)
    border_paragraph(h, color=GOLD, sz=10, sides=('bottom',), space=6)
    paragraph_spacing(h, before=0, after=14)
    add_toc(doc)
    note = doc.add_paragraph()
    rtl_paragraph(note)
    paragraph_spacing(note, before=10, after=0)
    add_text(note, "ملاحظة: يُحدَّث هذا الفهرس تلقائيًا من عناوين المستند "
                    "(حتى المستوى الثالث). عند فتح الملف في Word اضغط Ctrl+A "
                    "ثم F9 لعرضه.", name=FONT_BODY, size=9, color=GRAY_MED)


def chapter_divider(doc, kicker, title, teaser, eyebrow=None):
    for _ in range(3):
        doc.add_paragraph()
    if eyebrow:
        pe = doc.add_paragraph()
        rtl_paragraph(pe, align='right')
        paragraph_spacing(pe, before=0, after=2)
        add_text(pe, eyebrow, name=FONT_HEAD, size=10.5, bold=True, color=BROWN)
    pk = doc.add_paragraph()
    rtl_paragraph(pk, align='right')
    paragraph_spacing(pk, before=0, after=6)
    add_text(pk, kicker, name=FONT_HEAD, size=12, bold=True, color=GOLD)

    pt = doc.add_paragraph()
    rtl_paragraph(pt, align='right')
    border_paragraph(pt, color=GREEN, sz=16, sides=('right',), space=14)
    paragraph_indent(pt, left=0.3)
    paragraph_spacing(pt, before=0, after=14, line=1.2)
    add_text(pt, title, name=FONT_HEAD, size=26, bold=True, color=GREEN)

    if teaser:
        pz = doc.add_paragraph()
        rtl_paragraph(pz, align='right')
        paragraph_indent(pz, left=0.3)
        paragraph_spacing(pz, before=0, after=0, line=1.4)
        add_text(pz, teaser, name=FONT_BODY, size=12.5, color=GRAY_TEXT)
    _blank(doc, 130)
    add_picture_centered(doc, BANNER, width_cm=13.5)
    add_page_break(doc)


def render_heading(doc, level, text):
    text = clean(text)
    h = doc.add_heading(level=level)
    rtl_paragraph(h, align='right')
    r = h.add_run(text)
    set_font(r, name=FONT_HEAD)
    return h


def fix_run_boundaries(runs):
    """Insert a space between adjacent runs where one ends in an Arabic
    letter and the next starts with a digit (source authoring glitch,
    e.g. two runs 'الاستراتيجية' + '20%' concatenated with no space)."""
    out = []
    prev_last = ""
    for r in runs:
        t = r["t"]
        if t and prev_last and re.match(r'[ء-ي]', prev_last) and re.match(r'[0-9]', t):
            t = " " + t
        if t:
            prev_last = t[-1]
        out.append({**r, "t": t})
    return out


def render_paragraph(doc, b, bold_all=False):
    p = doc.add_paragraph()
    rtl_paragraph(p)
    paragraph_spacing(p, before=0, after=8, line=1.32)
    runs = b["runs"] if b["runs"] else [{"t": b["text"], "b": bold_all, "i": False, "u": False}]
    runs = fix_run_boundaries(runs)
    for r in runs:
        if r["t"] == "\n":
            p.add_run().add_break()
            continue
        add_text(p, r["t"], name=FONT_BODY, size=11, bold=r["b"], italic=r["i"],
                 underline=r["u"], color="242424")
    return p


def render_bullet(doc, b):
    lines = [clean(x) for x in b["text"].split("\n")]
    p = doc.add_paragraph(style='List Bullet')
    rtl_paragraph(p, align='right')
    paragraph_spacing(p, before=0, after=5, line=1.28)
    paragraph_indent(p, left=0.9)
    bold = lines[0]
    add_text(p, bold, name=FONT_BODY, size=11, bold=(len(lines) > 1), color="242424")
    for extra in lines[1:]:
        p2 = doc.add_paragraph()
        rtl_paragraph(p2, align='right')
        paragraph_indent(p2, left=0.9)
        paragraph_spacing(p2, before=0, after=5, line=1.25)
        add_text(p2, clean(extra), name=FONT_BODY, size=10, color=GRAY_TEXT)
    return p


def render_table(doc, b, order_idx):
    rows, cols = b["nrows"], b["ncols"]
    tbl = doc.add_table(rows=rows, cols=cols)
    for ri, row_vals in enumerate(b["rows"]):
        for ci, val in enumerate(row_vals):
            tbl.rows[ri].cells[ci].text = val
    style = TABLE_STYLE_BY_ORDER.get(order_idx, 'default')
    width = 15.5 / cols
    widths = [width] * cols
    if style == 'scale':
        style_scale_table(tbl, col_widths=widths)
    elif style == 'cf_scale':
        style_scale_table(tbl, col_widths=widths)
    elif style == 'role2col':
        style_role_table(tbl, col_widths=widths)
    elif style == 'weight3':
        style_table_default(tbl, col_widths=widths)
    elif style == 'weight2':
        style_table_default(tbl, col_widths=widths)
    elif style == 'log':
        style_log_table(tbl, col_widths=widths, extra_rows=3)
    else:
        style_table_default(tbl, col_widths=widths)
    spacer(doc)
    return tbl


def render_flow_paragraph(doc, text):
    orient, steps = split_flow(text)
    if orient == 'v':
        process_flow(doc, steps)
    else:
        process_flow_horizontal(doc, steps)


def render_example_range(doc, start, end, title_text=None):
    first = B[start]
    title = title_text or clean(first["text"])
    lines = []
    conclusion = None
    for i in range(start + 1, end + 1):
        blk = B[i]
        if blk["kind"] != "p":
            continue
        t = blk["text"].strip()
        if not t:
            continue
        if is_flow_text(t):
            _, steps = split_flow(t)
            lines.append("  ←  ".join(reversed(steps)))
            continue
        for sub in t.split("\n"):
            sub = sub.strip()
            if not sub:
                continue
            if sub.startswith("التوصية"):
                conclusion = sub
            else:
                lines.append(sub)
    example_box(doc, title, lines, conclusion=conclusion)


# =========================================================== main flow ==

def new_section(doc, chapter_title):
    from docx.enum.section import WD_SECTION
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    setup_section(sec)
    build_header(sec, f"{DOC_TITLE} · {chapter_title}" if chapter_title else DOC_TITLE)
    build_footer(sec, DOC_TITLE)
    return sec


def build():
    doc = Document()
    setup_styles(doc)
    sec0 = doc.sections[0]
    setup_section(sec0)  # cover page: page size/RTL only, no header/footer

    cover_page(doc)
    new_section(doc, None)
    toc_page(doc)

    table_order = -1
    i = 0
    front_eyebrow_text = clean(B[FRONT_EYEBROW_IDX]["text"]) if FRONT_EYEBROW_IDX in B else None
    current_chapter_label = "مقدمة الدليل"

    while i < N:
        b = B[i]

        if i < FRONT_EYEBROW_IDX:
            # original cover-page lines (0-1) and the hand-typed TOC (2-38):
            # both are fully superseded, verbatim, by cover_page()/toc_page().
            i += 1
            continue

        if i == FRONT_EYEBROW_IDX:
            i += 1
            continue

        if i in CHAPTER_STARTS:
            kicker, teaser_idx = CHAPTER_STARTS[i]
            title = clean(b["text"])
            teaser_full = B[teaser_idx]["text"]
            teaser = teaser_full.split("،")[0].split(".")[0].strip()
            if len(teaser) > 170:
                teaser = teaser[:170].rsplit(" ", 1)[0] + "…"
            eyebrow = front_eyebrow_text if i == 52 else None
            new_section(doc, title)
            chapter_divider(doc, kicker, title, teaser, eyebrow=eyebrow)
            current_chapter_label = title
            i += 1
            continue

        if b["kind"] == "tbl":
            table_order += 1
            render_table(doc, b, table_order)
            i += 1
            continue

        # ---- paragraph ----
        text = b["text"]

        if i in CASE_START_IDX:
            for s, e in CASE_RANGES:
                if s == i:
                    render_example_range(doc, s, e)
                    i = e + 1
                    break
            continue
        if i in CASE_SKIP:
            i += 1
            continue

        if i == SIMPLE_EXAMPLE_RANGE[0]:
            render_example_range(doc, *SIMPLE_EXAMPLE_RANGE, title_text="مثال مبسط")
            i = SIMPLE_EXAMPLE_RANGE[1] + 1
            continue
        if SIMPLE_EXAMPLE_RANGE[0] < i <= SIMPLE_EXAMPLE_RANGE[1]:
            i += 1
            continue

        if i == CLASS_LEGEND_TRIGGER:
            entries = []
            colors = [GREEN, "4f8a63", "c98b4c", RED]
            letters = ["A", "B", "C", "D"]
            for k, idx2 in enumerate((263, 264, 265, 266)):
                t = B[idx2]["text"]
                parts = t.split("\n")
                label = parts[0].strip()
                desc = parts[1].strip() if len(parts) > 1 else ""
                entries.append((letters[k], label, desc, colors[k]))
            classification_cards(doc, entries)
            i += 1
            continue
        if i in CLASS_LEGEND_SKIP:
            i += 1
            continue

        if i == FLOW_TRIGGER:
            steps, subs = [], []
            for idx2 in (279, 281, 283, 285, 287, 289):
                t = B[idx2]["text"]
                parts = t.split("\n")
                steps.append(parts[0].strip())
                subs.append(parts[1].strip() if len(parts) > 1 else None)
            process_flow(doc, steps, sub=subs, compact=True)
            i += 1
            continue
        if i in FLOW_SKIP:
            i += 1
            continue

        if not text.strip():
            i += 1
            continue

        if is_flow_text(text):
            render_flow_paragraph(doc, text)
            i += 1
            continue

        level = None
        if i in FORCE_HEADING:
            level = FORCE_HEADING[i]
        elif i in HEADING_LEVEL_FIX:
            level = HEADING_LEVEL_FIX[i]
        elif b["style"] in ("Heading 1", "Heading 2", "Heading 3"):
            level = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3}[b["style"]]

        if level:
            title_clean = clean(text)
            if title_clean in DEFINITION_HEADINGS:
                nxt = B.get(i + 1)
                if nxt and nxt["kind"] == "p" and nxt["text"].strip().startswith("يقصد"):
                    render_heading(doc, min(level, 4), title_clean)
                    i += 2
                    body = B[i]["text"] if i < N else ""
                    definition_box(doc, "", clean(body))
                    i += 1
                    continue
            render_heading(doc, min(level, 4), title_clean)
            i += 1
            continue

        if i not in PRINCIPLE_EXCLUDE and is_principle_paragraph(b):
            principle_card(doc, clean(text))
            i += 1
            continue

        if b["is_list"]:
            render_bullet(doc, b)
        else:
            render_paragraph(doc, b)

        if i in GATE_INSERTS:
            kind, items = GATE_INSERTS[i]
            if kind == 'flow':
                process_flow(doc, items, compact=True)
            else:
                concept_cards_grid(doc, items)

        if i == ESCALATION_AFTER_IDX:
            pt = doc.add_paragraph()
            rtl_paragraph(pt)
            paragraph_spacing(pt, before=10, after=4)
            add_text(pt, "مسار التصعيد", name=FONT_HEAD, size=12, bold=True, color=GREEN)
            process_flow(doc, [
                "الجهة المالكة للمؤشر أو المشروع",
                "فريق الاستراتيجية في الجهة",
                "مكتب إدارة الاستراتيجية (SMO)",
                "الجهة المخولة باتخاذ القرار — للحالات الاستثنائية",
            ], compact=True)

        i += 1

    force_update_fields(doc)
    doc.save(OUT_PATH)
    print("Saved", OUT_PATH)


if __name__ == "__main__":
    build()
